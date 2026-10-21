"""
Reports — Services
Business logic for report queryset scoping, date filtering, Excel/CSV generation,
and AuditLog tracking.
"""
import csv
from datetime import datetime
from io import BytesIO, StringIO
import logging
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from django.db.models import Count, Q, Sum
from django.utils import timezone
from rest_framework.exceptions import ValidationError

from apps.audit.models import AuditLog
from apps.payments.models import Payment
from apps.requests.models import Request
from apps.services.models import Service
from apps.workflow.constants import RequestStatuses

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Queryset Scoping & Filtering
# ---------------------------------------------------------------------------

def parse_date_param(param_value, field_name):
    """Parse YYYY-MM-DD string into a date object or raise ValidationError."""
    if not param_value:
        return None
    try:
        return datetime.strptime(param_value.strip(), "%Y-%m-%d").date()
    except ValueError:
        raise ValidationError({
            "error": {
                "code": "VALIDATION_ERROR",
                "message": "One or more query parameters are invalid.",
                "details": {
                    field_name: ["Date has wrong format. Use YYYY-MM-DD."]
                }
            }
        })


def get_scoped_request_queryset(user):
    """
    Returns the base non-draft Request queryset scoped to the user's role.
    - SYSTEM_ADMIN: all non-draft requests across the institution.
    - SERVICE_DEPT_ADMIN / STAFF: restricted to requests in their service department.
    """
    qs = Request.objects.exclude(current_status=RequestStatuses.DRAFT).select_related(
        'student_id',
        'student_id__academic_department_id',
        'service_id',
        'service_id__service_department_id',
        'current_step_id',
    )

    if not user.is_system_admin():
        dept = getattr(user, 'service_department_id', None)
        dept_id = getattr(dept, 'id', None) if dept else None
        if not dept_id:
            return Request.objects.none()
        qs = qs.filter(service_id__service_department_id=dept_id)

    return qs


def apply_request_filters(qs, params, user):
    """
    Applies standard filters: from_date, to_date, status, service_department_id,
    service_id, academic_department_id, stream, payment_status, search.
    """
    from_date = parse_date_param(params.get('from_date'), 'from_date')
    to_date = parse_date_param(params.get('to_date'), 'to_date')

    if from_date:
        qs = qs.filter(submitted_at__date__gte=from_date)
    if to_date:
        qs = qs.filter(submitted_at__date__lte=to_date)

    # Department filter (System Admin can filter by any department; Staff is locked to their own)
    dept_id = params.get('service_department_id')
    if dept_id and str(dept_id).isdigit():
        if user.is_system_admin():
            qs = qs.filter(service_id__service_department_id=int(dept_id))

    # Service filter
    service_id = params.get('service_id')
    if service_id and str(service_id).isdigit():
        qs = qs.filter(service_id=int(service_id))

    # Academic department filter
    acad_dept_id = params.get('academic_department_id')
    if acad_dept_id and str(acad_dept_id).isdigit():
        qs = qs.filter(student_id__academic_department_id=int(acad_dept_id))

    # Stream filter
    stream = params.get('stream')
    if stream and stream.strip().upper() in ['SFM', 'SFW', 'AIDED']:
        qs = qs.filter(student_id__stream__iexact=stream.strip())

    # Status filter
    status_param = params.get('status')
    if status_param and status_param.strip().upper() != 'ALL':
        qs = qs.filter(current_status__iexact=status_param.strip())

    # Payment status filter
    payment_status = params.get('payment_status')
    if payment_status:
        qs = qs.filter(payment__payment_status__iexact=payment_status.strip())

    # Search filter
    search = params.get('search', '').strip()
    if search:
        qs = qs.filter(
            Q(request_number__icontains=search) |
            Q(student_id__name__icontains=search) |
            Q(student_id__register_number__icontains=search)
        )

    # Ordering
    ordering = params.get('ordering', '-submitted_at').strip()
    valid_orderings = [
        'submitted_at', '-submitted_at',
        'completed_at', '-completed_at',
        'request_number', '-request_number',
        'current_status', '-current_status',
    ]
    if ordering in valid_orderings:
        qs = qs.order_by(ordering)
    else:
        qs = qs.order_by('-submitted_at')

    return qs


# ---------------------------------------------------------------------------
# Revenue Aggregation Service
# ---------------------------------------------------------------------------

def calculate_revenue_report(user, params):
    """
    Calculates revenue figures filtered by paid_at date range.
    Only payments with payment_status = 'SUCCESS' are aggregated.
    """
    from_date = parse_date_param(params.get('from_date'), 'from_date')
    to_date = parse_date_param(params.get('to_date'), 'to_date')

    # Payments with SUCCESS status linked to non-draft requests
    payments_qs = Payment.objects.filter(
        payment_status='SUCCESS',
        request_id__current_status__isnull=False,
    ).exclude(
        request_id__current_status=RequestStatuses.DRAFT
    ).select_related(
        'request_id',
        'request_id__service_id',
        'request_id__service_id__service_department_id',
    )

    # Scoping for department users
    if not user.is_system_admin():
        dept = getattr(user, 'service_department_id', None)
        dept_id = getattr(dept, 'id', None) if dept else None
        if not dept_id:
            payments_qs = Payment.objects.none()
        else:
            payments_qs = payments_qs.filter(request_id__service_id__service_department_id=dept_id)
    else:
        dept_id = params.get('service_department_id')
        if dept_id and str(dept_id).isdigit():
            payments_qs = payments_qs.filter(request_id__service_id__service_department_id=int(dept_id))

    # Payment date filtering on paid_at
    if from_date:
        payments_qs = payments_qs.filter(paid_at__date__gte=from_date)
    if to_date:
        payments_qs = payments_qs.filter(paid_at__date__lte=to_date)

    service_id = params.get('service_id')
    if service_id and str(service_id).isdigit():
        payments_qs = payments_qs.filter(request_id__service_id=int(service_id))

    payment_method = params.get('payment_method')
    if payment_method:
        payments_qs = payments_qs.filter(payment_method__iexact=payment_method.strip())

    # Overall Summary
    total_agg = payments_qs.aggregate(
        total_amount=Sum('amount'),
        successful_count=Count('id')
    )
    total_revenue_collected = f"{total_agg['total_amount'] or 0:.2f}"
    total_successful_transactions = total_agg['successful_count'] or 0

    # Group by Service
    by_service_rows = (
        payments_qs.values(
            'request_id__service_id',
            'request_id__service_id__code',
            'request_id__service_id__name',
            'request_id__service_id__service_department_id__name',
        )
        .annotate(
            successful_count=Count('id'),
            total_amount=Sum('amount'),
        )
        .order_by('-total_amount')
    )

    by_service = [
        {
            "service_id": row['request_id__service_id'],
            "service_code": row['request_id__service_id__code'],
            "service_name": row['request_id__service_id__name'],
            "service_department_name": row['request_id__service_id__service_department_id__name'] or '',
            "successful_count": row['successful_count'],
            "total_amount": f"{row['total_amount'] or 0:.2f}",
        }
        for row in by_service_rows
    ]

    # Group by Payment Method
    by_method_rows = (
        payments_qs.values('payment_method')
        .annotate(
            transaction_count=Count('id'),
            total_amount=Sum('amount'),
        )
        .order_by('-total_amount')
    )

    by_payment_method = [
        {
            "payment_method": row['payment_method'] or 'UNKNOWN',
            "transaction_count": row['transaction_count'],
            "total_amount": f"{row['total_amount'] or 0:.2f}",
        }
        for row in by_method_rows
    ]

    return {
        "filters_applied": {
            "from_date": params.get('from_date') or None,
            "to_date": params.get('to_date') or None,
            "service_department_id": int(dept_id) if dept_id and str(dept_id).isdigit() else None,
            "service_id": int(service_id) if service_id and str(service_id).isdigit() else None,
            "payment_method": payment_method or None,
        },
        "summary": {
            "total_revenue_collected": total_revenue_collected,
            "total_successful_transactions": total_successful_transactions,
            "currency": "INR",
        },
        "by_service": by_service,
        "by_payment_method": by_payment_method,
    }


# ---------------------------------------------------------------------------
# Excel (.xlsx) & CSV Export Generators
# ---------------------------------------------------------------------------

REQUEST_EXPORT_COLUMNS = [
    "Request Number",
    "Register Number",
    "Student Name",
    "Academic Department",
    "Stream",
    "Service Code",
    "Service Name",
    "Service Department",
    "Status",
    "Current Step",
    "Submitted On",
    "Completed On",
    "Fee (INR)",
    "Payment Status",
    "Payment Method",
    "Transaction ID",
]

REVENUE_EXPORT_COLUMNS = [
    "Service Code",
    "Service Name",
    "Service Department",
    "Successful Transactions",
    "Total Revenue Collected (INR)",
]


def generate_requests_data_rows(queryset):
    """Extract flat rows from Request queryset for export."""
    # Pre-fetch payments for speed
    payments_map = {
        p.request_id_id: p
        for p in Payment.objects.filter(request_id__in=queryset)
    }

    rows = []
    for req in queryset:
        student = getattr(req, 'student_id', None)
        acad_dept = getattr(student, 'academic_department_id', None) if student else None
        service = getattr(req, 'service_id', None)
        svc_dept = getattr(service, 'service_department_id', None) if service else None
        step = getattr(req, 'current_step_id', None)
        payment = payments_map.get(req.id)

        rows.append([
            req.request_number or '',
            student.register_number if student else '',
            student.name if student else '',
            acad_dept.name if acad_dept else '',
            student.stream if student else '',
            service.code if service else '',
            service.name if service else '',
            svc_dept.name if svc_dept else '',
            req.current_status or '',
            step.step_name if step else '',
            req.submitted_at.strftime('%Y-%m-%d %H:%M:%S') if req.submitted_at else '',
            req.completed_at.strftime('%Y-%m-%d %H:%M:%S') if req.completed_at else '',
            f"{payment.amount:.2f}" if payment and payment.amount is not None else '',
            payment.payment_status if payment else '',
            payment.payment_method if payment else '',
            payment.gateway_txn_id if payment else '',
        ])
    return rows


def generate_excel_export(headers, rows, sheet_name="Report"):
    """
    Builds an styled Excel workbook via openpyxl and returns bytes.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.append(headers)

    # Style header row
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # Append data
    for row in rows:
        ws.append(row)

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def generate_csv_export(headers, rows):
    """
    Generates UTF-8 CSV content and returns string.
    """
    buffer = StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    return buffer.getvalue()


def log_report_export(request, report_type, export_format, filters_dict):
    """Logs report export action into AuditLog."""
    try:
        AuditLog.log(
            request=request,
            action='EXPORT',
            app_label='reports',
            model_name='Request' if report_type == 'requests' else 'Payment',
            object_repr=f"Exported {report_type} report as {export_format.upper()}",
            changes={'filters': filters_dict, 'format': export_format, 'report_type': report_type}
        )
    except Exception as e:
        logger.error(f"Failed to record export AuditLog: {e}")
